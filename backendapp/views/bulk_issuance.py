import csv
import io
import uuid

from django.db import transaction
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from rest_framework import status

from backendapp.models import Document, DocumentType


def _generate_tracking_field():
    """Return an 8-char uppercase alphanumeric tracking code."""
    return str(uuid.uuid4()).replace('-', '').upper()[:8]


def _parse_csv(file_obj):
    """Parse an in-memory CSV file into a list of dicts (header → value)."""
    text = file_obj.read().decode('utf-8-sig')  # utf-8-sig strips BOM if present
    reader = csv.DictReader(io.StringIO(text))
    rows = []
    errors = []
    for i, row in enumerate(reader, start=2):  # row 1 is the header
        # Strip whitespace from keys and values
        cleaned = {k.strip(): v.strip() for k, v in row.items() if k}
        rows.append({'row': i, 'data': cleaned})
    return rows, errors


def _parse_excel(file_obj):
    """Parse an uploaded Excel (.xlsx) file into a list of dicts."""
    try:
        import openpyxl
    except ImportError:
        raise ImportError("openpyxl is required for Excel processing. Run: pip install openpyxl")

    wb = openpyxl.load_workbook(filename=io.BytesIO(file_obj.read()), data_only=True)
    ws = wb.active
    rows = []
    errors = []

    headers = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            headers = [str(h).strip() if h is not None else '' for h in row]
            continue
        row_data = {headers[j]: (str(v).strip() if v is not None else '') for j, v in enumerate(row)}
        rows.append({'row': i + 1, 'data': row_data})

    return rows, errors


class BulkIssuanceView(APIView):
    """
    POST /api/bulk-issuance/

    Accepts:
        - file          (CSV or .xlsx, multipart)
        - document_type (int, the DocumentType PK)

    Returns:
        {
          "created": 42,
          "errors": [{"row": 5, "reason": "..."}],
          "tracking_codes": ["AB23CD45", ...]
        }
    """
    permission_classes = [IsAuthenticated]

    def post(self, request):
        uploaded_file = request.FILES.get('file')
        document_type_id = request.data.get('document_type')

        # --- Validate inputs ---
        if not uploaded_file:
            return Response({'error': 'No file uploaded.'}, status=status.HTTP_400_BAD_REQUEST)

        if not document_type_id:
            return Response({'error': 'document_type is required.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            doc_type = DocumentType.objects.get(pk=int(document_type_id))
        except (DocumentType.DoesNotExist, ValueError):
            return Response({'error': 'Invalid document_type ID.'}, status=status.HTTP_400_BAD_REQUEST)

        # --- Parse the file ---
        filename = uploaded_file.name.lower()
        parse_errors = []
        try:
            if filename.endswith('.csv'):
                rows, parse_errors = _parse_csv(uploaded_file)
            elif filename.endswith(('.xlsx', '.xls')):
                rows, parse_errors = _parse_excel(uploaded_file)
            else:
                return Response(
                    {'error': 'Unsupported file type. Please upload a .csv or .xlsx file.'},
                    status=status.HTTP_400_BAD_REQUEST,
                )
        except ImportError as e:
            return Response({'error': str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)
        except Exception as e:
            return Response({'error': f'Failed to parse file: {str(e)}'}, status=status.HTTP_400_BAD_REQUEST)

        if not rows:
            return Response({'error': 'The uploaded file contains no data rows.'}, status=status.HTTP_400_BAD_REQUEST)

        # --- Build Document objects ---
        documents_to_create = []
        row_errors = list(parse_errors)
        tracking_codes = []

        for row_info in rows:
            row_num = row_info['row']
            row_data = row_info['data']

            # Skip entirely empty rows
            if not any(row_data.values()):
                continue

            tracking_code = _generate_tracking_field()
            tracking_codes.append(tracking_code)

            documents_to_create.append(
                Document(
                    document_type=doc_type,
                    tracking_field=tracking_code,
                    metadata=row_data,
                    issued_by=request.user,
                    status='valid',
                )
            )

        # --- Count BEFORE bulk_create (safe on SQLite which may not return IDs) ---
        success_count = len(documents_to_create)

        if success_count == 0:
            return Response(
                {'error': 'No valid rows found in the file.', 'errors': row_errors},
                status=status.HTTP_400_BAD_REQUEST,
            )

        # --- Persist all at once ---
        with transaction.atomic():
            Document.objects.bulk_create(documents_to_create)

        return Response({
            'created': success_count,
            'errors': row_errors,
            'tracking_codes': tracking_codes,
        }, status=status.HTTP_201_CREATED)
